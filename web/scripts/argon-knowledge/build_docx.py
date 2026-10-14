import json, re, sys, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SP = sys.argv[1]
OUT = sys.argv[2]
kb = json.load(open(f"{SP}/creator_kb.json", encoding="utf-8"))
wk = json.load(open(f"{SP}/kb_wiki_pages.json", encoding="utf-8"))
KB = {r["title"]: r for r in kb}
WK = {r["title"]: r for r in wk}

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
TODAY = "14/10/2026"

# ───────── cấu trúc tài liệu: (mã phần, tiêu đề phần, mô tả, [(nguồn, tiêu đề gốc, tiêu đề hiển thị)]) ─────────
PARTS = [
 ("1", "Tổng quan GoHub & thuật ngữ", "GoHub là ai, bán gì, đo hiệu quả bằng chỉ số nào.", [
   ("wiki", "GoHub — Tổng Quan Công Ty", None),
   ("wiki", "Thuật Ngữ Chỉ Số Kinh Doanh (Business Metrics)", None),
 ]),
 ("2", "Sản phẩm: mã SKU, Item, Combo, Data Policy", "Cách đặt mã và phân loại sản phẩm. Đây là nền để đọc đúng mọi mã SKU/Item trong dữ liệu.", [
   ("wiki", "Cấu Trúc Mã SKU (13 ký tự)", None),
   ("wiki", "Cấu Trúc Mã Item & Alias (18–20 ký tự)", None),
   ("wiki", "Data Policy Codes (Mã Loại Data)", None),
   ("wiki", "Combo Chuẩn GoHub (42 combo/country)", None),
   ("kb", "Định nghĩa Trạng thái SKU: Temporary", None),
   ("kb", "Quy chuẩn mã sản phẩm và Combo 42 gói", None),
   ("kb", "Quy tắc kiểm tra trạng thái mở bán Ecom (Mã DVE tổng vs Mã DES/DET lẻ)", None),
   ("kb", "Chi Phí Phôi eSIM Profile và Công Thức Tính eSIM Full", None),
 ]),
 ("3", "Nhà cung cấp (Vendor)", "Thứ tự ưu tiên chọn vendor, hồ sơ từng vendor, chính sách đổi/hủy/hoàn tiền.", [
   ("wiki", "Chọn Vendor Nào? — Hướng Dẫn Nhanh", None),
   ("wiki", "Vendor 3HK — Hướng Dẫn Tư Vấn", None),
   ("wiki", "Vendor WorldMove (WM) — Hướng Dẫn Tư Vấn", None),
   ("wiki", "Chính Sách Vendor: QR, Đổi Thiết Bị, Hủy/Hoàn Tiền", None),
   ("kb", "Phân biệt Vendor BC (Billion Connect) và WD (BC Datapool)", None),
   ("kb", "Chính sách & Quy định sản phẩm Vendor GigHub (Airhub)", None),
   ("kb", "Chính sách & Quy định sản phẩm Vendor Viettech", None),
   ("kb", "Báo giá & Đánh giá Vendor Roam Communication (T8/2026)", None),
   ("kb", "Cơ chế Top-up Data cộng dồn (Add-on) 3HK Datapool", None),
 ]),
 ("4", "Giá vốn (COGS), tỷ giá & báo giá", "Công thức tính giá nhập, giả định tiêu dùng, tỷ giá nội bộ và các bảng báo giá pool.", [
   ("live", "Giả định & tham số giá vốn HIỆN HÀNH", None),
   ("wiki", "Công Thức Tính Giá Nhập — 3HK", None),
   ("wiki", "Định Giá Gói FUP Tùy Chỉnh 3HK", None),
   ("wiki", "Tỷ Giá Nội Bộ GoHub", None),
   ("kb", "Quy tắc tính giá COGS chung cho các Vendor Datapool (3HK, BC Datapool, Vendor mới)", None),
   ("kb", "Cơ Chế Phí & Thanh Toán BC Datapool (Singtel & CMHK)", None),
   ("kb", "Báo Giá BC Datapool - CMHK & Singtel Pool (23/09/2026)", None),
   ("kb", "Bảng giá và vùng phủ BC Datapool (Singtel & CMHK Core) - 23/09/2026", None),
   ("kb", "Bảng Sản Phẩm & Giá Gốc VNPT (Inbound VN & Outbound 16 Nước)", None),
   ("kb", "Quy tắc nhà cung cấp và công thức tính COGS", None),
   ("kb", "3HK Datapool Assumptions for COGS Calculation", "3HK Datapool — giả định COGS (bản ghi cũ, tháng 08/2026)"),
   ("kb", "Quy tắc Nguồn Tỷ giá Chuẩn từ Supabase", None),
   ("kb", "Tỷ giá hiện tại (gohub_dw)", None),
 ]),
 ("5", "Quy trình CS & vận hành", "Các quy trình làm việc hằng ngày của CS, Ops và Product.", [
   ("wiki", "Quy Trình CS & Vận Hành: SIMPIN, Tồn Kho Cận Date, Báo Cáo Tuần", None),
   ("wiki", "Quy Trình Import NCC — Upload → Kiểm tra → Xác nhận", None),
   ("kb", "Hướng dẫn check mã SIMPIN đối me với sản phẩm Gighub (Airhub) cho CS", "Hướng dẫn tra mã SIMPIN sản phẩm GigHub (Airhub) cho CS"),
   ("kb", "Playbook Quy Trình Xử Lý Tồn Kho & Phôi SIM Cận Date", None),
   ("kb", "Quy trình báo cáo tuần nhóm KH Non-Strategic", None),
   ("kb", "Chỉ số BI và quy trình vận hành Ops", None),
 ]),
 ("6", "Mục tiêu & KPI", "Lộ trình chỉ tiêu Q3–Q4/2026 và điểm xuất phát (ghi chú tháng 08/2026 — số liệu thực tế đã thay đổi, xem mục “Độ mới của thông tin”).", [
   ("kb", "Lộ Trình KPI & Targets Q3 - Q4 2026 — Hiếu", "Lộ trình KPI & Targets Q3–Q4/2026 (Product Ops & BI)"),
   ("kb", "Mục Tiêu & Điểm Xuất Phát Q3 (Tháng 8 - Tháng 10/2026) — Hiếu", "Mục tiêu & điểm xuất phát Q3/2026 (Product Ops & BI)"),
   ("kb", "Thông tin cá nhân & Mục tiêu KPI của Hiếu", "Vai trò Product Ops & BI và KPI cá nhân Q3–Q4/2026"),
 ]),
 ("7", "Ghi chú thị trường & khách hàng", "Khuyến mãi, tương thích ứng dụng theo nhà mạng và các ghi chú khách hàng.", [
   ("kb", "Chính sách khuyến mãi SIM/eSIM Thái Lan DTAC & TrueMove H (Hạn 31/12/2026)", None),
   ("kb", "Bảng Tương Thích Ứng Dụng (TikTok, AI, Maps) Theo Nhà Mạng & IP Gốc", None),
   ("kb", "Khách hàng B2B Strategic: OD US Sim Station", None),
   ("kb", "Template chung cho Wellzone và Infigate", None),
 ]),
]

LIVE_MD = """Phần này là **giá trị đang được cấu hình trong hệ thống** (Admin › Cài đặt › Công Thức Datapool và Tỷ Giá Nội Bộ), đo ngày 14/10/2026. **Khi các ghi chú cũ ở phần này hoặc ở các mục khác nói khác, ưu tiên phần này.**

### Giả định tiêu dùng data (để tính giá vốn Datapool)

| Loại gói | Giả định hiện hành | Cách tính data thực dùng |
|---|---|---|
| Fixed (dung lượng cố định) | 55% dung lượng gói | Dung lượng gói × 0,55 |
| Daily (theo ngày) | 35% dung lượng ngày | GB/ngày × số ngày × 0,35 |
| Unlimited 500MB + 5Mbps | 1,9 GB/ngày | 1,9 × số ngày |
| Unlimited 500MB + 10Mbps | 2,0 GB/ngày | 2,0 × số ngày |
| Unlimited 3GB + 10Mbps | 2,0 GB/ngày | 2,0 × số ngày |

Bảng mức Unlimited có thể thêm mức mới (GB tốc độ cao, mức throttle Mbps, GB/ngày) ở Công Thức Datapool; mức nào chưa có trong bảng thì hệ thống không đoán giá.

### Tỷ giá nội bộ tháng 10/2026

| Tỷ giá | Giá trị | Dùng để |
|---|---|---|
| HKD / USD (Inc) | 7,807 | Đổi giá pool tính bằng HKD sang USD |
| CNY / USD (Inc) | 6,681 | Đổi phí eSIM tính bằng CNY sang USD |
| VND / USD (JSC) | 26.153 | Đổi giá vốn USD sang VND (đầu bán VN) |
| VND / USD (Inc) | 25.626,2 | Đổi giá VND sang USD |

### Quy tắc làm tròn và công thức (chốt 14/10/2026)

- **Làm tròn MỘT lần ở tiền cuối, luôn làm tròn lên (ROUNDUP):** USD làm tròn 2 số lẻ; VND làm tròn 0 số lẻ. Không làm tròn trung gian ở bước HKD.
- Giá data (USD) = ROUNDUP(giá pool theo GB × GB thực dùng [÷ HKD/USD], 2).
- eSIM full = ROUNDUP(data USD + phí khung eSIM, 2); đầu VN = ROUNDUP(giá USD × VND/USD JSC, 0).
- **SIM full BC (đầu VN, không có đầu US):** VND = ROUNDUP(data USD × VND/USD JSC, 0) + giá SIM trắng (VND). Không cộng phí IMSI và không đổi giá SIM trắng qua USD.
- **SIM full 3HK:** đầu VN = ROUNDUP(data USD × VND/USD JSC, 0) + giá khung SIM VN (VND).
- Giá gói data rời (datapack, mã A) chỉ gồm data, không cộng phí khung.

### Vendor Datapool trong hệ thống giá vốn

- **3HK Datapool:** mã vendor 3D; giá pool theo nước tính bằng HKD/GB.
- **BC Datapool CMHK:** mã vendor WD; giá pool bằng HKD/GB.
- **BC Datapool Singtel:** mã vendor W1; giá pool bằng USD/GB.
- Hệ thống có tab **Soát giá vốn** tính lại mọi SKU 3HK/BC theo các tham số trên và nhắn DM Lark khi giá đang lưu lệch giá tính lại.
"""

EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿️⭐⬆✅❌]")
CRED = re.compile(r"^\s*(?:[-*]|\d+\.)?\s*\**\s*(account|password|username|tài khoản|mật khẩu|otp)\b\**\s*:", re.I)


def sanitize(md: str) -> str:
    out = []
    dropped = False
    for ln in md.splitlines():
        if CRED.match(ln) and not re.search(r"otp", ln, re.I):
            dropped = True
            continue
        out.append(ln)
    t = "\n".join(out)
    if dropped:
        n = 0; res = []
        for ln in t.split("\n"):
            m = re.match(r"^(\d+)\.\s+(.*)$", ln)
            if m:
                n += 1; res.append(f"{n}. {m.group(2)}")
            else:
                if ln.strip() and not ln.startswith(" "): n = 0
                res.append(ln)
        t = "\n".join(res)
        t += "\n\n> Thông tin đăng nhập portal không ghi trong tài liệu này (bảo mật). Nhân viên cần tài khoản liên hệ quản lý để được cấp.\n"
    t = re.sub(r"[\w.]+@gohub\.com", "[tài khoản dùng chung — xin quản lý]", t)
    t = EMOJI.sub("", t)
    t = re.sub(r"[ \t]+\n", "\n", t)
    return t


# ───────── python-docx helpers ─────────
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)


def set_font(style, name="Calibri", size=None, bold=None, color=None, italic=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color


st = doc.styles
set_font(st["Normal"], size=10.5)
st["Normal"].paragraph_format.space_after = Pt(5)
st["Normal"].paragraph_format.line_spacing = 1.12
set_font(st["Title"], size=30, bold=True, color=NAVY)
set_font(st["Heading 1"], size=19, bold=True, color=NAVY)
st["Heading 1"].paragraph_format.space_before = Pt(0)
st["Heading 1"].paragraph_format.space_after = Pt(8)
set_font(st["Heading 2"], size=14.5, bold=True, color=NAVY)
st["Heading 2"].paragraph_format.space_before = Pt(16)
st["Heading 2"].paragraph_format.space_after = Pt(5)
set_font(st["Heading 3"], size=12, bold=True, color=RGBColor(0x2E, 0x5C, 0x8A))
st["Heading 3"].paragraph_format.space_before = Pt(11)
st["Heading 3"].paragraph_format.space_after = Pt(3)
set_font(st["Heading 4"], size=10.5, bold=True, color=RGBColor(0x40, 0x40, 0x40))
st["Heading 4"].paragraph_format.space_before = Pt(8)
st["Heading 4"].paragraph_format.space_after = Pt(2)
for n in ("List Bullet", "List Bullet 2", "List Bullet 3"):
    set_font(st[n], size=10.5)
    st[n].paragraph_format.space_after = Pt(2)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)


def borders(table, color="BFBFBF"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); trPr.append(e)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")


def fixmd(t):
    t = re.sub(r"\[\[([^\]|]+)\|([^\]]+)\]\]", r"\2", t)
    t = re.sub(r"\[\[([^\]]+)\]\]", r"\1", t)
    if t.count("**") % 2: t = t.replace("**", "")
    if t.count("`") % 2: t = t.replace("`", "")
    return t


def add_inline(p, text, size=None, color=None, bold=None):
    text = fixmd(text)
    text = text.replace("<br>", "\n").replace("<br/>", "\n").replace("&nbsp;", " ")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2].replace("`", "")); r.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt((size or 10.5) - 1)
            r.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1].replace("`", "")); r.italic = True
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            r = p.add_run(f"{m.group(1)} ({m.group(2)})") if m else p.add_run(part)
        else:
            r = p.add_run(part)
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
        if bold: r.bold = True


def add_table(rows):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    borders(t); cell_margins(t)
    fs = 9.5 if ncol <= 4 else 8.5 if ncol <= 6 else 7.5
    for i, r in enumerate(rows):
        row = t.rows[i]
        no_split(row)
        if i == 0: repeat_header(row)
        for j, c in enumerate(r):
            cell = row.cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i == 0:
                shade(cell, "1F3A5F")
                add_inline(p, c, size=fs, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
            else:
                if i % 2 == 0: shade(cell, "F2F6FA")
                add_inline(p, c, size=fs)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(ln):
    s = ln.strip()
    if s.startswith("|"): s = s[1:]
    if s.endswith("|"): s = s[:-1]
    return [c.strip() for c in s.split("|")]


def render_md(md, base_level=3):
    """base_level: cấp heading Word cho '##' trong nguồn. '#' đầu tiên (trùng tiêu đề mục) bị bỏ."""
    lines = md.splitlines()
    i = 0
    first_h1_skipped = False
    num_counter = {}
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1; continue
        if re.match(r"^-{3,}$|^\*{3,}$|^_{3,}$", s):
            i += 1; continue
        if s.startswith("```"):
            buf = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F3F3"); pPr.append(shd)
            for k, b in enumerate(buf):
                r = p.add_run(b)
                r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas"); r.font.size = Pt(8.5)
                if k < len(buf) - 1: r.add_break()
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1)); txt = m.group(2).strip().strip("#").strip()
            txt = re.sub(r"\*\*|`", "", txt)
            if lvl == 1 and not first_h1_skipped:
                first_h1_skipped = True; i += 1; continue
            if not txt:
                i += 1; continue
            h = min(4, base_level + max(0, lvl - 2)) if lvl >= 2 else base_level
            doc.add_heading(txt, level=h)
            i += 1; continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            rows = [split_row(s)]; i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            add_table(rows); continue
        if s.startswith(">"):
            buf = [s.lstrip("> ").strip()]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(\s*([-*•]|\d+[.)])\s|#{1,6}\s|\|)", lines[i]):
                buf.append(lines[i].strip().lstrip("> ").strip()); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, fixmd(" ".join(buf)), color=GREY)
            for r in p.runs: r.italic = True
            continue
        m = re.match(r"^(\s*)([-*•])\s+(.*)$", ln)
        if m:
            depth = min(2, len(m.group(1).replace("\t", "    ")) // 2 // 1 // 2) if m.group(1) else 0
            style = ["List Bullet", "List Bullet 2", "List Bullet 3"][depth]
            p = doc.add_paragraph(style=style)
            txt = m.group(3).strip(); i += 1
            while i < len(lines) and re.match(r"^\s{2,}\S", lines[i]) and not re.match(r"^\s*([-*•]|\d+[.)])\s|^\s*\|", lines[i]):
                txt += " " + lines[i].strip(); i += 1
            add_inline(p, fixmd(txt))
            continue
        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", ln)
        if m:
            depth = 1 if len(m.group(1)) >= 2 else 0
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9 + 0.7 * depth)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(f"{m.group(2)}.\t"); r.bold = True
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9 + 0.7 * depth))
            txt = m.group(3).strip(); i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^\s*([-*•]|\d+[.)])\s|^\s*\||^#{1,6}\s|^>|^```", lines[i]) and re.match(r"^\s{2,}\S|^\S", lines[i]) and lines[i].startswith(" "):
                txt += " " + lines[i].strip(); i += 1
            add_inline(p, fixmd(txt))
            continue
        # đoạn văn thường: gộp các dòng liền nhau
        buf = [s]; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(\s*([-*•]|\d+[.)])\s|#{1,6}\s|\||>|```|-{3,}$)", lines[i]):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph()
        add_inline(p, fixmd(" ".join(buf)))


def add_field(par, instr):
    r = par.add_run()
    for t, txt in (("begin", None), (None, instr), ("separate", None), (None, "Nhấn F9 để cập nhật mục lục"), ("end", None)):
        if t:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
        elif txt == instr:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r._r.append(it)
        else:
            tt = OxmlElement("w:t"); tt.text = txt; r._r.append(tt)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ───────── bìa ─────────
for _ in range(5): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("GOHUB"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x2E, 0x5C, 0x8A)
r.font.name = "Calibri"
t = doc.add_paragraph("Sổ tay Tri thức Công ty", style="Title")
p = doc.add_paragraph()
r = p.add_run("Nghiệp vụ SIM/eSIM · Vendor · Giá vốn · Quy trình vận hành · Mục tiêu KPI")
r.font.size = Pt(13); r.font.color.rgb = GREY
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Tài liệu tổng hợp để trợ lý AI (Gemini 4 Argon) đọc và trả lời câu hỏi nghiệp vụ.")
r.font.size = Pt(11)
doc.add_paragraph()
meta = doc.add_table(rows=4, cols=2)
borders(meta, "D9D9D9"); cell_margins(meta, 60, 60, 100, 100)
for k, (a, b) in enumerate([("Phiên bản", f"Tổng hợp ngày {TODAY}"), ("Nguồn", "Wiki GoHub (nghiệp vụ) và các ghi chú Gấu Pro đã học (Creator KB)"),
                            ("Phạm vi", "Chỉ kiến thức nghiệp vụ dùng chung. Không gồm thông tin đăng nhập, trí nhớ cá nhân hay dữ liệu khách hàng."),
                            ("Quản lý", "Hiếu — Product Operations & BI")]):
    meta.rows[k].cells[0].text = ""; meta.rows[k].cells[1].text = ""
    add_inline(meta.rows[k].cells[0].paragraphs[0], a, size=10, bold=True, color=NAVY)
    add_inline(meta.rows[k].cells[1].paragraphs[0], b, size=10)
    shade(meta.rows[k].cells[0], "EEF3F8")
    meta.rows[k].cells[0].width = Cm(3.4); meta.rows[k].cells[1].width = Cm(13)
page_break()

# ───────── hướng dẫn đọc ─────────
doc.add_heading("Hướng dẫn đọc tài liệu này", level=1)
render_md("""### Dành cho trợ lý AI

- Tài liệu chia **7 phần**; mỗi phần gồm nhiều mục. Trích dẫn nguồn theo **tên mục** khi trả lời (ví dụ: “theo mục *Chọn Vendor Nào?*”).
- **Thứ tự ưu tiên khi hai nơi nói khác nhau:** (1) mục *Giả định & tham số giá vốn HIỆN HÀNH* (Phần 4) → (2) trang Wiki nghiệp vụ → (3) ghi chú Gấu Pro đã học. Ghi chú có **ngày cập nhật** ở đầu mỗi mục; ghi chú cũ hơn dễ đã lỗi thời.
- Số liệu thị trường, giá báo của vendor và KPI **thay đổi theo tháng**. Với câu hỏi về giá/tỷ giá/chỉ số, luôn nêu ngày của nguồn và khuyến nghị kiểm tra lại dữ liệu sống trong hệ thống GoHub Intel.
- Không có thông tin đăng nhập portal trong tài liệu. Nếu được hỏi, trả lời rằng cần xin quản lý cấp.

### Thuật ngữ nhanh

| Từ | Nghĩa |
|---|---|
| SKU | Mã bán hàng 13 ký tự (xem Phần 2) |
| COGS | Giá vốn hàng bán (giá nhập từ vendor) |
| CM1 | Biên đóng góp cấp 1 = doanh thu − giá vốn − chi phí kênh/khách |
| Datapool | Gói data dùng chung theo pool của vendor, giá vốn tính theo dung lượng thực dùng |
| 3HK / BC / WM | Vendor 3HK (Hong Kong), BillionConnect, WorldMove |
| eSIM full / SIM full / datapack | Gói gồm SIM/eSIM + data / gói chỉ data nạp thêm (mã A) |
| B2B / B2C | Bán sỉ cho doanh nghiệp / bán lẻ |
""", base_level=2)
page_break()

# ───────── mục lục ─────────
doc.add_heading("Mục lục", level=1)
pt = doc.add_paragraph()
add_field(pt, 'TOC \\o "1-2" \\h \\z \\u')
page_break()

# ───────── nội dung ─────────
manifest = []
missing = []
for code, title, desc, items in PARTS:
    doc.add_heading(f"Phần {code}. {title}", level=1)
    p = doc.add_paragraph(); add_inline(p, desc, color=GREY)
    for r in p.runs: r.italic = True
    for k, (src, key, disp) in enumerate(items, 1):
        if src == "live":
            doc.add_heading(f"{code}.{k} {key}", level=2)
            meta_p = doc.add_paragraph(); add_inline(meta_p, "Nguồn: cấu hình hệ thống · Cập nhật: 14/10/2026", size=9, color=GREY)
            render_md(LIVE_MD, base_level=3)
            manifest.append((code, key, "Cấu hình hệ thống", "14/10/2026"))
            continue
        rec = (WK if src == "wiki" else KB).get(key)
        if rec is None:
            missing.append(key); continue
        name = disp or key
        doc.add_heading(f"{code}.{k} {name}", level=2)
        label = "Wiki GoHub" if src == "wiki" else "Ghi chú Gấu Pro (Creator KB)"
        upd = (rec.get("updated_at") or "")[:10]
        upd_vn = "/".join(reversed(upd.split("-"))) if upd else "?"
        mp = doc.add_paragraph(); add_inline(mp, f"Nguồn: {label} · Cập nhật: {upd_vn}", size=9, color=GREY)
        body = sanitize(rec["content"] or "")
        render_md(body, base_level=3)
        manifest.append((code, name, label, upd_vn))
    if code != PARTS[-1][0]:
        page_break()

# ───────── phụ lục ─────────
page_break()
doc.add_heading("Phụ lục. Độ mới của thông tin & danh mục nguồn", level=1)
render_md("""### Lưu ý về độ mới

- Các ghi chú Gấu Pro tháng 08/2026 (ví dụ *3HK Datapool — giả định COGS*, *Lộ trình KPI*) phản ánh **thời điểm đó**. Từ đó giả định giá vốn đã đổi (Daily từ 38% xuống 35%; Unlimited đổi mức GB/ngày) — dùng mục “Giả định & tham số giá vốn HIỆN HÀNH” ở Phần 4.
- Bảng báo giá BC Datapool (CMHK, Singtel) ghi theo bản **23/09/2026**; vendor có thể đổi giá hằng tháng.
- Mục tiêu KPI Q3–Q4/2026 là kế hoạch, không phải kết quả thực tế.

### Danh mục nguồn đã đưa vào tài liệu
""", base_level=2)
add_table([["Phần", "Mục", "Nguồn", "Cập nhật"]] + [[a, b, c, d] for a, b, c, d in manifest])
render_md("""### Nội dung không đưa vào tài liệu

- Thông tin đăng nhập portal vendor (tài khoản, mật khẩu): bỏ vì bảo mật.
- Trí nhớ cá nhân của từng người dùng với trợ lý AI.
- Hướng dẫn kỹ thuật từng tab của hệ thống GoHub Intel và tài liệu kiến trúc hệ thống: đã có trong Wiki hệ thống, không phải kiến thức nghiệp vụ.
- Dữ liệu doanh thu/khách hàng thực tế: lấy từ hệ thống GoHub Intel (có tab “Gói dữ liệu” riêng trong Phòng Argon).
""", base_level=2)

# ───────── header/footer ─────────
sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]
hp.text = ""
r = hp.add_run("GoHub — Sổ tay Tri thức Công ty"); r.font.size = Pt(8.5); r.font.color.rgb = GREY
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Trang "); r.font.size = Pt(8.5); r.font.color.rgb = GREY
r2 = fp.add_run()
for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if t:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r2._r.append(e)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r2._r.append(it)
r2.font.size = Pt(8.5); r2.font.color.rgb = GREY

# cập nhật mục lục khi mở
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)
doc.core_properties.title = "GoHub — Sổ tay Tri thức Công ty"
doc.core_properties.author = "GoHub Intel"
doc.core_properties.subject = "Tri thức nghiệp vụ cho Gemini 4 Argon"
doc.save(OUT)
print("saved", OUT, "items", len(manifest), "missing", missing)
